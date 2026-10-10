from io import BytesIO
from tempfile import TemporaryDirectory
from datetime import timedelta
from unittest.mock import patch
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from openpyxl import Workbook
from PIL import Image
from rest_framework.test import APITestCase
from rest_framework import status
from rest_framework_simplejwt.tokens import AccessToken, RefreshToken
from google.genai.errors import ServerError
import hashlib
import hmac
import uuid
from unittest.mock import Mock
from .models import Store, Membership, Product, Category, Order, OrderItem

class TenantIsolationTests(APITestCase):
    def setUp(self):
        User=get_user_model()
        self.owner=User.objects.create_user(username='one',email='one@example.com',password='long-test-password')
        self.other=User.objects.create_user(username='two',email='two@example.com',password='long-test-password')
        self.store=Store.objects.create(owner=self.owner,name='One Shop',slug='one-shop')
        Membership.objects.create(store=self.store,user=self.owner,role='owner')
        self.product=Product.objects.create(store=self.store,name='Book',price='10.00',stock=3)

    def test_anonymous_cannot_list_private_products(self):
        response=self.client.get('/api/stores/one-shop/products/')
        self.assertEqual(response.status_code,status.HTTP_401_UNAUTHORIZED)

    def test_nonmember_cannot_read_store_products(self):
        self.client.force_authenticate(self.other)
        response=self.client.get('/api/stores/one-shop/products/')
        self.assertEqual(response.status_code,status.HTTP_404_NOT_FOUND)

    def test_member_reads_only_own_store_catalog(self):
        self.client.force_authenticate(self.owner)
        response=self.client.get('/api/stores/one-shop/products/')
        self.assertEqual(response.status_code,status.HTTP_200_OK)
        self.assertEqual([row['id'] for row in response.data],[self.product.id])

    def test_jwt_login_and_vite_cors_preflight(self):
        response=self.client.post('/api/auth/login/', {'email':self.owner.email,'password':'long-test-password'}, format='json')
        self.assertEqual(response.status_code,status.HTTP_200_OK)
        self.assertIn('access',response.data); self.assertIn('refresh',response.data)
        self.assertEqual(self.client.get('/api/auth/me/',HTTP_AUTHORIZATION=f"Bearer {response.data['access']}").status_code,status.HTTP_200_OK)
        stores=self.client.get('/api/stores/',HTTP_AUTHORIZATION=f"Bearer {response.data['access']}")
        self.assertEqual(stores.status_code,status.HTTP_200_OK)
        self.assertEqual([row['slug'] for row in stores.data['results']],['one-shop'])
        preflight=self.client.options('/api/auth/login/',HTTP_ORIGIN='http://localhost:5174',HTTP_ACCESS_CONTROL_REQUEST_METHOD='POST',HTTP_ACCESS_CONTROL_REQUEST_HEADERS='content-type')
        self.assertEqual(preflight.status_code,status.HTTP_200_OK)
        self.assertEqual(preflight['Access-Control-Allow-Origin'],'http://localhost:5174')
        loopback=self.client.options('/api/stores/',HTTP_ORIGIN='http://127.0.0.1:5173',HTTP_ACCESS_CONTROL_REQUEST_METHOD='GET',HTTP_ACCESS_CONTROL_REQUEST_HEADERS='authorization')
        self.assertEqual(loopback['Access-Control-Allow-Origin'],'http://127.0.0.1:5173')

    def test_invalid_and_expired_bearer_tokens_are_rejected(self):
        invalid=self.client.get('/api/stores/',HTTP_AUTHORIZATION='Bearer not-a-jwt')
        self.assertEqual(invalid.status_code,status.HTTP_401_UNAUTHORIZED)
        expired=AccessToken.for_user(self.owner)
        expired.set_exp(lifetime=timedelta(seconds=-1))
        response=self.client.get('/api/stores/',HTTP_AUTHORIZATION=f'Bearer {expired}')
        self.assertEqual(response.status_code,status.HTTP_401_UNAUTHORIZED)

    def test_refresh_token_returns_access_accepted_by_protected_routes(self):
        refresh=str(RefreshToken.for_user(self.owner))
        response=self.client.post('/api/auth/refresh/',{'refresh':refresh},format='json')
        self.assertEqual(response.status_code,status.HTTP_200_OK)
        self.assertIn('access',response.data)
        stores=self.client.get('/api/stores/',HTTP_AUTHORIZATION=f"Bearer {response.data['access']}")
        self.assertEqual(stores.status_code,status.HTTP_200_OK)

    def test_store_creation_requires_jwt(self):
        response=self.client.post('/api/stores/',{'name':'Private Shop','slug':'private-shop'},format='json')
        self.assertEqual(response.status_code,status.HTTP_401_UNAUTHORIZED)

    def test_public_store_requires_published_flag(self):
        hidden=self.client.get('/api/public/stores/one-shop/')
        self.assertEqual(hidden.status_code,status.HTTP_404_NOT_FOUND)
        self.store.is_published=True; self.store.save(update_fields=['is_published'])
        public=self.client.get('/api/public/stores/one-shop/')
        self.assertEqual(public.status_code,status.HTTP_200_OK)
        self.assertNotIn('owner',public.data)

    def test_store_create_and_private_route_require_authenticated_owner(self):
        login=self.client.post('/api/auth/login/',{'email':self.owner.email,'password':'long-test-password'},format='json')
        token=login.data['access']
        created=self.client.post('/api/stores/',{'name':'New Shop','slug':'new-shop'},format='json',HTTP_AUTHORIZATION=f'Bearer {token}')
        self.assertEqual(created.status_code,status.HTTP_201_CREATED)
        self.assertEqual(created.data['slug'],'new-shop')

    def test_store_name_generates_unique_slug_and_preserves_existing_store(self):
        self.client.force_authenticate(self.owner)
        first=self.client.post('/api/stores/',{'name':'Sweet Store'},format='json')
        duplicate=self.client.post('/api/stores/',{'name':'Sweet Store'},format='json')
        self.assertEqual(first.status_code,status.HTTP_201_CREATED)
        self.assertEqual(duplicate.status_code,status.HTTP_201_CREATED)
        self.assertEqual(first.data['slug'],'sweet-store')
        self.assertEqual(duplicate.data['slug'],'sweet-store-2')
        self.store.refresh_from_db()
        self.assertEqual(self.store.slug,'one-shop')

    def test_public_store_does_not_include_products_from_other_tenants(self):
        self.store.is_published=True
        self.store.save(update_fields=['is_published'])
        self.product.sku='PRIVATE-SKU'
        self.product.low_stock_threshold=99
        self.product.save(update_fields=['sku','low_stock_threshold'])
        another=Store.objects.create(owner=self.other,name='Another',slug='another')
        foreign=Product.objects.create(store=another,name='Foreign item',price='3.00',stock=1)
        response=self.client.get('/api/public/stores/one-shop/')
        self.assertEqual(response.status_code,status.HTTP_200_OK)
        self.assertEqual([row['id'] for row in response.data['products']],[self.product.id])
        self.assertNotIn(foreign.id,[row['id'] for row in response.data['products']])
        self.assertNotIn('sku',response.data['products'][0])
        self.assertNotIn('low_stock_threshold',response.data['products'][0])
        self.assertNotIn('owner',response.data)

    def test_invalid_or_unpublished_public_store_url_returns_not_found(self):
        self.store.is_published=False
        self.store.save(update_fields=['is_published'])
        self.assertEqual(self.client.get('/api/public/stores/missing-store/').status_code,status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.get('/api/public/stores/one-shop/').status_code,status.HTTP_404_NOT_FOUND)

    def test_category_duplicate_returns_validation_error(self):
        self.client.force_authenticate(self.owner)
        first=self.client.post('/api/stores/one-shop/categories/',{'name':'Tea'},format='json')
        self.assertEqual(first.status_code,status.HTTP_201_CREATED)
        duplicate=self.client.post('/api/stores/one-shop/categories/',{'name':'tea'},format='json')
        self.assertEqual(duplicate.status_code,status.HTTP_400_BAD_REQUEST)

    def test_public_product_detail_only_serves_active_real_products(self):
        self.store.is_published=True; self.store.save(update_fields=['is_published'])
        response=self.client.get(f'/api/public/stores/one-shop/products/{self.product.id}/')
        self.assertEqual(response.status_code,status.HTTP_200_OK)
        demo=Product.objects.create(store=self.store,name='Demo',price='2.00',stock=1,is_demo=True)
        public_store=self.client.get('/api/public/stores/one-shop/')
        self.assertNotIn(demo.id,[item['id'] for item in public_store.data['products']])
        demo_detail=self.client.get(f'/api/public/stores/one-shop/products/{demo.id}/')
        self.assertEqual(demo_detail.status_code,status.HTTP_404_NOT_FOUND)

    def test_theme_and_publish_endpoints_persist_store_values(self):
        self.client.force_authenticate(self.owner)
        themes=self.client.get('/api/stores/one-shop/theme/')
        self.assertEqual(themes.status_code,status.HTTP_200_OK)
        updated=self.client.patch('/api/stores/one-shop/theme/',{'theme':'premium-dark','theme_settings':{'font_family':'Georgia, serif'}},format='json')
        self.assertEqual(updated.status_code,status.HTTP_200_OK)
        self.store.refresh_from_db()
        self.assertEqual(self.store.theme,'premium-dark')
        self.assertEqual(self.store.theme_settings['font_family'],'Georgia, serif')
        published=self.client.post('/api/stores/one-shop/publish/',{},format='json')
        self.assertEqual(published.status_code,status.HTTP_200_OK)
        self.store.refresh_from_db(); self.assertTrue(self.store.is_published)

    def test_product_create_rejects_zero_price_and_allows_repeated_blank_sku(self):
        self.client.force_authenticate(self.owner)
        zero=self.client.post('/api/stores/one-shop/products/',{'name':'Free','price':'0','stock':1},format='json')
        self.assertEqual(zero.status_code,status.HTTP_400_BAD_REQUEST)
        for name in ('No SKU 1','No SKU 2'):
            response=self.client.post('/api/stores/one-shop/products/',{'name':name,'sku':'','price':'2.00','stock':1},format='json')
            self.assertEqual(response.status_code,status.HTTP_201_CREATED)
        self.assertEqual(Product.objects.filter(store=self.store,name__startswith='No SKU').count(),2)

    def test_slug_availability_is_public_but_does_not_expose_store_data(self):
        response=self.client.get('/api/stores/slug-availability/',{'slug':'one-shop'})
        self.assertEqual(response.status_code,status.HTTP_200_OK)
        self.assertEqual(response.data,{'slug':'one-shop','available':False})

class InventoryAssistantTests(APITestCase):
    def setUp(self):
        self.owner=get_user_model().objects.create_user(username='owner',email='owner@example.com',password='long-test-password')
        self.other=get_user_model().objects.create_user(username='other',email='other@example.com',password='long-test-password')
        self.store=Store.objects.create(owner=self.owner,name='Inventory',slug='inventory')
        Membership.objects.create(store=self.store,user=self.owner,role='owner')
        self.category=Category.objects.create(store=self.store,name='Tea')
        self.product=Product.objects.create(store=self.store,category=self.category,name='Green Tea',sku='TEA-1',price='4.50',stock=2,low_stock_threshold=5)
        self.client.force_authenticate(self.owner)

    @override_settings(GEMINI_API_KEY='test-key',GEMINI_MODEL='test-model')
    @patch('apps.stores.assistant_views.answer_inventory_question', return_value='Green Tea has 2 units.')
    def test_assistant_uses_scoped_inventory_and_returns_evidence(self, mocked):
        response=self.client.post('/api/ai/chat/',{'store_slug':'inventory','question':'What is low? '},format='json')
        self.assertEqual(response.status_code,status.HTTP_200_OK)
        self.assertEqual(response.data['evidence']['product_count'],1)
        self.assertEqual(response.data['evidence']['low_stock'][0]['name'],'Green Tea')
        mocked.assert_called_once()

    @override_settings(GEMINI_API_KEY='test-key',GEMINI_MODEL='test-model')
    @patch('apps.stores.assistant_views.answer_inventory_question', return_value='One real order is present.')
    def test_assistant_evidence_contains_store_scoped_order_aggregates(self, mocked):
        order=Order.objects.create(store=self.store,status=Order.Status.DELIVERED,total='9.00')
        OrderItem.objects.create(order=order,product=self.product,name='Green Tea',quantity=2,unit_price='4.50')
        response=self.client.post('/api/ai/chat/',{'store_slug':'inventory','question':'How many orders today?'},format='json')
        self.assertEqual(response.status_code,status.HTTP_200_OK)
        facts=mocked.call_args.args[1]
        self.assertEqual(facts['sales_reports']['today']['order_count'],1)
        self.assertEqual(facts['sales_reports']['today']['order_value'],'9.00')
        self.assertEqual(facts['sales_reports']['today']['supporting_orders'][0]['id'],order.id)
        self.assertEqual(facts['sales_reports']['this_month']['top_products'][0]['units'],2)
        self.assertIn('end_date_exclusive',facts['sales_reports']['this_month'])
        self.assertEqual(facts['sales_reports']['currency'],'not configured on the store')

    @override_settings(GEMINI_API_KEY='',GEMINI_MODEL='')
    def test_assistant_reports_missing_configuration(self):
        response=self.client.post('/api/ai/chat/',{'store_slug':'inventory','question':'How much stock?'},format='json')
        self.assertEqual(response.status_code,status.HTTP_503_SERVICE_UNAVAILABLE)

    @override_settings(GEMINI_API_KEY='test-key',GEMINI_MODEL='test-model')
    @patch('apps.stores.assistant_views.answer_inventory_question', side_effect=Exception('upstream'))
    def test_assistant_handles_provider_failure_without_leaking_error(self, mocked):
        response=self.client.post('/api/ai/chat/',{'store_slug':'inventory','question':'How much stock?'},format='json')
        self.assertEqual(response.status_code,status.HTTP_502_BAD_GATEWAY)
        self.assertNotIn('upstream',str(response.data))

    @override_settings(GEMINI_API_KEY='test-key',GEMINI_MODEL='test-model')
    @patch('apps.stores.assistant_views.answer_inventory_question', side_effect=__import__('httpx').ConnectError('secret request detail'))
    def test_assistant_reports_connectivity_failure_without_upstream_details(self, mocked):
        response=self.client.post('/api/ai/chat/',{'store_slug':'inventory','question':'How much stock?'},format='json')
        self.assertEqual(response.status_code,status.HTTP_502_BAD_GATEWAY)
        self.assertIn('could not connect to Gemini',response.data['detail'])
        self.assertNotIn('secret request detail',str(response.data))

    @override_settings(GEMINI_API_KEY='test-key',GEMINI_MODEL='test-model')
    @patch('apps.stores.assistant_views.answer_inventory_question', side_effect=__import__('httpx').ReadTimeout('sensitive timeout detail'))
    def test_assistant_reports_timeout_as_gateway_timeout(self, mocked):
        response=self.client.post('/api/ai/chat/',{'store_slug':'inventory','question':'How much stock?'},format='json')
        self.assertEqual(response.status_code,status.HTTP_504_GATEWAY_TIMEOUT)
        self.assertIn('15 second timeout',response.data['detail'])
        self.assertNotIn('sensitive timeout detail',str(response.data))

    @override_settings(GEMINI_API_KEY='test-key',GEMINI_MODEL='test-model')
    @patch('apps.stores.assistant_views.answer_inventory_question', side_effect=ServerError(504, {'error': {'message': 'private provider detail'}}, None))
    def test_assistant_preserves_gemini_deadline_status(self, mocked):
        response=self.client.post('/api/ai/chat/',{'store_slug':'inventory','question':'How much stock?'},format='json')
        self.assertEqual(response.status_code,status.HTTP_504_GATEWAY_TIMEOUT)
        self.assertIn('exceeded its response deadline',response.data['detail'])
        self.assertNotIn('private provider detail',str(response.data))

    @override_settings(GEMINI_API_KEY='test-key',GEMINI_MODEL='test-model')
    @patch('apps.stores.assistant_views.answer_inventory_question', return_value='Answer')
    def test_assistant_cannot_read_another_store(self, mocked):
        self.client.force_authenticate(self.other)
        response=self.client.post('/api/ai/chat/',{'store_slug':'inventory','question':'How much stock?'},format='json')
        self.assertEqual(response.status_code,status.HTTP_404_NOT_FOUND)
        mocked.assert_not_called()

class ProductImportTests(APITestCase):
    def setUp(self):
        self.owner=get_user_model().objects.create_user(username='importer',email='importer@example.com',password='long-test-password')
        self.other=get_user_model().objects.create_user(username='outsider',email='outsider@example.com',password='long-test-password')
        self.store=Store.objects.create(owner=self.owner,name='Importer Shop',slug='importer-shop')
        Membership.objects.create(store=self.store,user=self.owner,role='owner')
        self.category=Category.objects.create(store=self.store,name='Tea')
        self.client.force_authenticate(self.owner)

    def upload(self, contents, name='products.csv', commit='false', slug='importer-shop'):
        file=SimpleUploadedFile(name,contents,content_type='text/csv')
        return self.client.post(f'/api/stores/{slug}/products/import/',{'file':file,'commit':commit},format='multipart')

    def test_csv_preview_does_not_write_and_confirm_import_persists(self):
        csv=b'name,sku,description,price,stock,category,image_url\nGreen Tea,TEA-1,Fresh,4.50,12,Tea,\n'
        preview=self.upload(csv)
        self.assertEqual(preview.status_code,status.HTTP_200_OK)
        self.assertTrue(preview.data['preview_only'])
        self.assertEqual(Product.objects.filter(store=self.store).count(),0)
        saved=self.upload(csv,commit='true')
        self.assertEqual(saved.status_code,status.HTTP_200_OK)
        self.assertEqual(saved.data['successful_rows'],1)
        self.assertEqual(Product.objects.get(store=self.store).category,self.category)

    def test_missing_required_columns_report_header_row(self):
        response=self.upload(b'name,price\nTea,4.50\n')
        self.assertEqual(response.status_code,status.HTTP_207_MULTI_STATUS)
        self.assertEqual(response.data['errors'][0]['row'],1)
        self.assertEqual(Product.objects.filter(store=self.store).count(),0)

    def test_invalid_price_reports_csv_row(self):
        response=self.upload(b'name,price,stock,category\nTea,nope,4,Tea\n')
        self.assertEqual(response.data['errors'][0]['row'],2)
        self.assertIn('price',response.data['errors'][0]['errors'][0])

    def test_negative_stock_reports_csv_row(self):
        response=self.upload(b'name,price,stock,category\nTea,4.50,-1,Tea\n')
        self.assertEqual(response.data['errors'][0]['row'],2)
        self.assertTrue(any('stock' in error for error in response.data['errors'][0]['errors']))

    def test_unknown_category_is_rejected(self):
        response=self.upload(b'name,price,stock,category\nTea,4.50,4,Coffee\n')
        self.assertEqual(response.data['errors'][0]['row'],2)
        self.assertIn('category',response.data['errors'][0]['errors'][0])

    def test_duplicate_sku_is_reported_without_importing(self):
        Product.objects.create(store=self.store,name='Existing',sku='SAME',price='2.00',stock=1)
        response=self.upload(b'name,sku,price,stock,category\nAnother,SAME,3.00,4,Tea\n')
        self.assertEqual(response.data['failed_rows'],1)
        self.assertIn('SKU already exists',response.data['errors'][0]['errors'][0])

    def test_empty_file_reports_header_row(self):
        response=self.upload(b'')
        self.assertEqual(response.data['errors'][0]['row'],1)
        self.assertIn('empty',response.data['errors'][0]['errors'][0])

    def test_user_without_store_membership_cannot_import(self):
        self.client.force_authenticate(self.other)
        response=self.upload(b'name,price,stock,category\nTea,4.50,4,Tea\n')
        self.assertEqual(response.status_code,status.HTTP_404_NOT_FOUND)

    def test_category_is_required_and_sku_is_optional(self):
        response=self.upload(b'name,price,stock,category\nTea,4.50,4,Tea\n',commit='true')
        self.assertEqual(response.status_code,status.HTTP_200_OK)
        product=Product.objects.get(store=self.store)
        self.assertIsNone(product.sku)

    def test_xlsx_is_parsed_as_values_only(self):
        wb=Workbook(); ws=wb.active; ws.append(['name','price','stock','category']); ws.append(['Mint',2.5,8,'Tea'])
        stream=BytesIO(); wb.save(stream)
        response=self.upload(stream.getvalue(),name='products.xlsx')
        self.assertEqual(response.status_code,status.HTTP_200_OK)
        self.assertEqual(response.data['valid_rows'],1)

    def test_xlsx_template_is_store_authorized_and_uses_existing_category(self):
        response=self.client.get('/api/stores/importer-shop/products/import/template/')
        self.assertEqual(response.status_code,status.HTTP_200_OK)
        self.assertIn('spreadsheetml.sheet',response['Content-Type'])
        from openpyxl import load_workbook
        generated=load_workbook(BytesIO(response.content),read_only=True,data_only=True)
        self.assertEqual(list(next(generated.active.iter_rows(values_only=True))),['name','price','stock','category','sku','description','image_url'])
        self.assertEqual(generated.active['D2'].value,'Tea')
        self.client.force_authenticate(self.other)
        denied=self.client.get('/api/stores/importer-shop/products/import/template/')
        self.assertEqual(denied.status_code,status.HTTP_404_NOT_FOUND)

    def test_owner_can_upload_store_logo_image(self):
        image=BytesIO(); Image.new('RGB',(2,2),'green').save(image,format='PNG')
        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            upload=SimpleUploadedFile('logo.png',image.getvalue(),content_type='image/png')
            response=self.client.post(f'/api/stores/{self.store.id}/assets/',{'logo_image':upload},format='multipart')
            self.assertEqual(response.status_code,status.HTTP_200_OK)
            self.assertIn('/media/stores/logos/',response.data['logo_image'])
            self.client.force_authenticate(self.other)
            denied_upload=SimpleUploadedFile('logo.png',image.getvalue(),content_type='image/png')
            denied=self.client.post(f'/api/stores/{self.store.id}/assets/',{'logo_image':denied_upload},format='multipart')
            self.assertEqual(denied.status_code,status.HTTP_404_NOT_FOUND)


class CheckoutOrderTests(APITestCase):
    def setUp(self):
        User=get_user_model()
        self.customer=User.objects.create_user(username='buyer',email='buyer@example.com',password='long-test-password')
        self.owner=User.objects.create_user(username='seller',email='seller@example.com',password='long-test-password')
        self.outsider=User.objects.create_user(username='outsider',email='outsider@example.com',password='long-test-password')
        self.store=Store.objects.create(owner=self.owner,name='Published Shop',slug='published-shop',is_published=True)
        Membership.objects.create(store=self.store,user=self.owner,role='owner')
        self.product=Product.objects.create(store=self.store,name='Lamp',price='12.50',stock=8)
        self.client.force_authenticate(self.customer)

    def checkout(self, *, key=None, method='cod', product_id=None, quantity=2, extra=None):
        key=key or str(uuid.uuid4())
        body={
            'customer_name':'Buyer', 'customer_phone':'+919876543210',
            'delivery_address':{'line1':'1 Test Road','city':'Pune','postal_code':'411001','country':'India'},
            'payment_method':method,
            'items':[{'product_id':product_id or self.product.id,'quantity':quantity}],
        }
        body.update(extra or {})
        return self.client.post('/api/public/stores/published-shop/checkout/',body,format='json',HTTP_IDEMPOTENCY_KEY=key),key

    def test_cod_checkout_uses_database_price_and_reserves_stock_idempotently(self):
        response,key=self.checkout(extra={'total':'0.01','unit_price':'0.01'})
        self.assertEqual(response.status_code,status.HTTP_201_CREATED)
        self.assertEqual(response.data['subtotal'],'25.00')
        self.assertEqual(response.data['total'],'25.00')
        self.assertEqual(response.data['payment_method'],'cod')
        self.assertEqual(response.data['payment_status'],'pending')
        self.assertTrue(response.data['order_number'].startswith('SS-'))
        self.product.refresh_from_db(); self.assertEqual(self.product.stock,6)
        again,_=self.checkout(key=key,extra={'total':'9999.00'})
        self.assertEqual(again.status_code,status.HTTP_200_OK)
        self.assertEqual(again.data['order_number'],response.data['order_number'])
        self.assertEqual(Order.objects.count(),1)
        self.product.refresh_from_db(); self.assertEqual(self.product.stock,6)
        changed,_=self.checkout(key=key,quantity=1)
        self.assertEqual(changed.status_code,status.HTTP_409_CONFLICT)
        self.assertEqual(Order.objects.count(),1)

    def test_checkout_requires_authentication(self):
        self.client.force_authenticate(user=None)
        response,_=self.checkout()
        self.assertEqual(response.status_code,status.HTTP_401_UNAUTHORIZED)

    def test_insufficient_stock_and_foreign_store_product_create_no_order(self):
        response,_=self.checkout(quantity=99)
        self.assertEqual(response.status_code,status.HTTP_409_CONFLICT)
        self.assertEqual(Order.objects.count(),0)
        foreign=Store.objects.create(owner=self.outsider,name='Other',slug='other',is_published=True)
        other_product=Product.objects.create(store=foreign,name='Foreign',price='1.00',stock=10)
        response,_=self.checkout(product_id=other_product.id)
        self.assertEqual(response.status_code,status.HTTP_400_BAD_REQUEST)
        self.assertEqual(Order.objects.count(),0)

    @override_settings(RAZORPAY_MODE='test',RAZORPAY_KEY_ID='',RAZORPAY_KEY_SECRET='')
    def test_missing_gateway_config_keeps_cod_but_rejects_card_without_order(self):
        response,_=self.checkout(method='razorpay')
        self.assertEqual(response.status_code,status.HTTP_503_SERVICE_UNAVAILABLE)
        self.assertEqual(Order.objects.count(),0)
        cod,_=self.checkout()
        self.assertEqual(cod.status_code,status.HTTP_201_CREATED)

    @override_settings(RAZORPAY_MODE='test',RAZORPAY_KEY_ID='rzp_test_example',RAZORPAY_KEY_SECRET='test-secret')
    @patch('apps.stores.order_views.httpx.post')
    def test_razorpay_order_uses_server_total_and_signature_verification(self, gateway_post):
        provider=Mock(); provider.json.return_value={'id':'order_provider_1'}; provider.raise_for_status.return_value=None
        gateway_post.return_value=provider
        checkout,_=self.checkout(method='razorpay')
        self.assertEqual(checkout.status_code,status.HTTP_201_CREATED)
        order_no=checkout.data['order_number']
        created=self.client.post(f'/api/orders/{order_no}/payment/create/',{},format='json')
        self.assertEqual(created.status_code,status.HTTP_201_CREATED)
        self.assertEqual(created.data['amount'],2500)
        self.assertEqual(gateway_post.call_args.kwargs['json']['amount'],2500)

        signature=hmac.new(b'test-secret',b'order_provider_1|pay_provider_1',hashlib.sha256).hexdigest()
        provider_payment=Mock(); provider_payment.json.return_value={
            'order_id':'order_provider_1','amount':2500,'currency':'INR','status':'captured'
        }; provider_payment.raise_for_status.return_value=None
        with patch('apps.stores.order_views.httpx.get',return_value=provider_payment) as gateway_get:
            verified=self.client.post(f'/api/orders/{order_no}/payment/verify/',{
                'razorpay_order_id':'order_provider_1','razorpay_payment_id':'pay_provider_1','razorpay_signature':signature,
            },format='json')
            self.assertEqual(verified.status_code,status.HTTP_200_OK)
            self.assertEqual(verified.data['payment_status'],'paid')
            self.assertEqual(gateway_get.call_count,1)
            duplicate=self.client.post(f'/api/orders/{order_no}/payment/verify/',{
                'razorpay_order_id':'order_provider_1','razorpay_payment_id':'pay_provider_1','razorpay_signature':signature,
            },format='json')
            self.assertEqual(duplicate.status_code,status.HTTP_200_OK)
            self.assertEqual(gateway_get.call_count,1)

    @override_settings(RAZORPAY_MODE='test',RAZORPAY_KEY_ID='rzp_test_example',RAZORPAY_KEY_SECRET='test-secret')
    @patch('apps.stores.order_views.httpx.get')
    def test_invalid_payment_signature_never_marks_paid(self, gateway_get):
        checkout,_=self.checkout(method='razorpay')
        order_no=checkout.data['order_number']
        provider=Mock(); provider.json.return_value={'id':'order_provider_2'}; provider.raise_for_status.return_value=None
        with patch('apps.stores.order_views.httpx.post',return_value=provider):
            self.client.post(f'/api/orders/{order_no}/payment/create/',{},format='json')
        response=self.client.post(f'/api/orders/{order_no}/payment/verify/',{
            'razorpay_order_id':'order_provider_2','razorpay_payment_id':'pay_fake','razorpay_signature':'invalid',
        },format='json')
        self.assertEqual(response.status_code,status.HTTP_400_BAD_REQUEST)
        gateway_get.assert_not_called()
        order=Order.objects.get(order_number=order_no)
        self.assertEqual(order.payment_status,Order.PaymentStatus.PENDING)

    @override_settings(RAZORPAY_MODE='test',RAZORPAY_KEY_ID='rzp_test_example',RAZORPAY_KEY_SECRET='test-secret')
    @patch('apps.stores.order_views.httpx.post',side_effect=__import__('httpx').ConnectError('private provider data'))
    def test_gateway_failure_marks_payment_failed_and_releases_stock(self, gateway_post):
        checkout,_=self.checkout(method='razorpay')
        self.assertEqual(checkout.status_code,status.HTTP_201_CREATED)
        failed=self.client.post(f"/api/orders/{checkout.data['order_number']}/payment/create/",{},format='json')
        self.assertEqual(failed.status_code,status.HTTP_502_BAD_GATEWAY)
        order=Order.objects.get(order_number=checkout.data['order_number'])
        self.assertEqual(order.payment_status,Order.PaymentStatus.FAILED)
        self.product.refresh_from_db(); self.assertEqual(self.product.stock,8)
        self.assertNotIn('private provider data',str(failed.data))

    @override_settings(RAZORPAY_MODE='test',RAZORPAY_KEY_ID='rzp_test_example',RAZORPAY_KEY_SECRET='test-secret')
    def test_customer_payment_failure_callback_releases_stock_once(self):
        checkout,_=self.checkout(method='razorpay')
        order_no=checkout.data['order_number']
        provider=Mock(); provider.json.return_value={'id':'order_provider_fail'}; provider.raise_for_status.return_value=None
        with patch('apps.stores.order_views.httpx.post',return_value=provider):
            self.client.post(f'/api/orders/{order_no}/payment/create/',{},format='json')
        first=self.client.post(f'/api/orders/{order_no}/payment/fail/',{},format='json')
        second=self.client.post(f'/api/orders/{order_no}/payment/fail/',{},format='json')
        self.assertEqual(first.status_code,status.HTTP_200_OK)
        self.assertEqual(second.status_code,status.HTTP_200_OK)
        self.product.refresh_from_db(); self.assertEqual(self.product.stock,8)
        self.assertEqual(Order.objects.get(order_number=order_no).payment_status,Order.PaymentStatus.FAILED)

    def test_customer_history_is_private_and_store_status_transitions_are_validated(self):
        created,_=self.checkout()
        number=created.data['order_number']
        self.client.force_authenticate(self.outsider)
        self.assertEqual(self.client.get('/api/stores/published-shop/orders/').status_code,status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.get(f'/api/orders/my/{number}/').status_code,status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.get('/api/orders/my/').data,[])
        self.client.force_authenticate(self.owner)
        list_response=self.client.get('/api/stores/published-shop/orders/')
        self.assertEqual(list_response.status_code,status.HTTP_200_OK)
        valid=self.client.patch(f'/api/stores/published-shop/orders/{number}/status/',{'status':'confirmed'},format='json')
        self.assertEqual(valid.status_code,status.HTTP_200_OK)
        invalid=self.client.patch(f'/api/stores/published-shop/orders/{number}/status/',{'status':'delivered'},format='json')
        self.assertEqual(invalid.status_code,status.HTTP_409_CONFLICT)
        self.assertEqual(Order.objects.get(order_number=number).status_history.count(),2)

    def test_store_owner_can_progress_order_and_confirm_delivered_cod_collection(self):
        created,_=self.checkout()
        number=created.data['order_number']
        self.client.force_authenticate(self.owner)
        for next_status in ('confirmed','processing','shipped','delivered'):
            response=self.client.patch(f'/api/stores/published-shop/orders/{number}/status/',{'status':next_status},format='json')
            self.assertEqual(response.status_code,status.HTTP_200_OK)
        collected=self.client.post(f'/api/stores/published-shop/orders/{number}/payment/collect/',{},format='json')
        self.assertEqual(collected.status_code,status.HTTP_200_OK)
        self.assertEqual(collected.data['payment_status'],'paid')
        self.assertEqual(collected.data['status'],'delivered')

    def test_customer_cancellation_releases_reserved_inventory_once(self):
        created,_=self.checkout()
        url=f"/api/orders/my/{created.data['order_number']}/cancel/"
        cancelled=self.client.post(url,{},format='json')
        self.assertEqual(cancelled.status_code,status.HTTP_200_OK)
        self.product.refresh_from_db(); self.assertEqual(self.product.stock,8)
        again=self.client.post(url,{},format='json')
        self.assertEqual(again.status_code,status.HTTP_409_CONFLICT)
        self.product.refresh_from_db(); self.assertEqual(self.product.stock,8)
